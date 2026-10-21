from __future__ import annotations

import io
import tempfile
import unittest
from pathlib import Path

import pymupdf as pdf
from docx import Document
from openpyxl import load_workbook
from PIL import Image
from pptx import Presentation

from office_export import export_office, page_layout
from settings import DEFAULTS


def formatted_pdf() -> bytes:
    with pdf.open() as doc:
        page = doc.new_page(width=500, height=650)
        page.insert_text((50, 75), "Titulo editable", fontname="hebo", fontsize=20, color=(1, 0, 0))
        page.insert_text((50, 110), "Detalle fuera de tabla", fontname="heit", fontsize=12)
        page.draw_rect((50, 150, 350, 190), color=None, fill=(0.8, 0.9, 1))
        for y in (150, 190, 230):
            page.draw_line((50, y), (350, y))
        for x in (50, 350):
            page.draw_line((x, 150), (x, 230))
        page.draw_line((200, 190), (200, 230))
        page.insert_text((60, 175), "Cabecera combinada", fontname="hebo", fontsize=14)
        page.insert_text((60, 215), "Dato fuera de formula")
        page.insert_text((210, 215), "=1+1", fontname="hebo", color=(0, 0, 1))
        image = Image.new("RGB", (120, 60), "green")
        buffer = io.BytesIO()
        image.save(buffer, format="PNG")
        page.insert_image((50, 280, 170, 340), stream=buffer.getvalue())
        page.draw_rect((50, 400, 100, 430), fill=(1, 0.5, 0))
        return doc.tobytes()


def word_background_image(document) -> bytes:
    images = [rel.target_part.blob for rel in document.part.rels.values() if rel.reltype.endswith("/image")]
    assert len(images) == 1, len(images)
    return images[0]


class OfficeExportTests(unittest.TestCase):
    def assert_rgb(self, actual: str, expected: str):
        for index in (0, 2, 4):
            self.assertAlmostEqual(int(actual[index:index + 2], 16),
                                   int(expected[index:index + 2], 16), delta=1)

    def setUp(self):
        self.temporary = tempfile.TemporaryDirectory()
        self.addCleanup(self.temporary.cleanup)
        self.folder = Path(self.temporary.name)
        self.doc = pdf.open(stream=formatted_pdf(), filetype="pdf")
        self.addCleanup(self.doc.close)
        self.config = DEFAULTS.copy()

    def convert(self, operation):
        return export_office(operation, self.doc, [0], self.folder, self.config)

    def test_layout_detects_native_merged_table_and_visuals(self):
        page = self.doc[0]
        layout = page_layout(page, self.config)
        self.assertFalse(layout["fallback"])
        self.assertEqual(len(layout["tables"]), 1)
        self.assertIn((0, 0, 0, 1), [cell[:4] for cell in layout["tables"][0]["cells"]])
        self.assertGreaterEqual(len(layout["visuals"]), 2)

    def test_word_styles_tables_images_and_page_dimensions(self):
        output, notice = self.convert("pdf_to_docx")
        document = Document(output)
        self.assertAlmostEqual(document.sections[0].page_width.pt, 500, places=1)
        self.assertAlmostEqual(document.sections[0].page_height.pt, 650, places=1)
        from docx.oxml.ns import qn
        paragraph = next(item for item in document.paragraphs if item.text == "Titulo editable")
        frame = paragraph._p.pPr.find(qn("w:framePr"))
        self.assertAlmostEqual(int(frame.get(qn("w:x"))) / 20, 50, delta=1)
        self.assertEqual(frame.get(qn("w:hAnchor")), "page")
        title = paragraph.runs[0]
        self.assertEqual(title.text, "Titulo")
        self.assertTrue(title.bold)
        self.assertEqual(title.font.size.pt, 20)
        self.assertEqual(str(title.font.color.rgb), "FF0000")
        self.assertTrue(any(run.italic for paragraph in document.paragraphs for run in paragraph.runs))
        self.assertEqual(len(document.tables), 1)
        table = document.tables[0]
        self.assertEqual(table.cell(0, 0).text, "Cabecera combinada")
        self.assertIs(table.cell(0, 0)._tc, table.cell(0, 1)._tc)
        self.assertIn("=1+1", table.cell(1, 1).text)
        position = table._tbl.tblPr.find(qn("w:tblpPr"))
        self.assertEqual((int(position.get(qn("w:tblpX"))), int(position.get(qn("w:tblpY")))), (1000, 3000))
        self.assertEqual(len(document.inline_shapes), 0)
        with Image.open(io.BytesIO(word_background_image(document))) as background:
            background = background.convert("RGB")
            self.assertEqual(background.size, (1000, 1300))
            # Fills, images and vectors stay in the background; rebuilt text ink does not.
            self.assert_rgb("".join(f"{value:02X}" for value in background.getpixel((600, 370))), "CCE6FF")
            self.assertEqual(background.getpixel((200, 600)), (0, 128, 0))
            pixels = background.crop((100, 110, 420, 155)).tobytes()
            self.assertEqual(sum(pixels[i] > 150 and pixels[i + 1] < 100 and pixels[i + 2] < 100
                                 for i in range(0, len(pixels), 3)), 0)
        self.assertIn("posicion original", notice)

    def test_word_rounds_fractional_font_sizes_instead_of_truncating(self):
        with pdf.open() as source:
            source.new_page().insert_text((50, 80), "Encabezado", fontname="hebo", fontsize=9.96)
            output, _ = export_office("pdf_to_docx", source, [0], self.folder, self.config)
        run = next(run for paragraph in Document(output).paragraphs for run in paragraph.runs
                   if run.text == "Encabezado")
        self.assertEqual(run.font.size.pt, 10)

    def test_excel_keeps_non_table_text_merges_styles_and_visual_reference(self):
        output, _ = self.convert("pdf_to_xlsx")
        workbook = load_workbook(output)
        self.addCleanup(workbook.close)
        self.assertEqual(workbook.sheetnames, ["Pagina 1", "Original 1"])
        sheet = workbook["Pagina 1"]
        cells = [cell for row in sheet for cell in row]
        title = next(cell for cell in cells if cell.value == "Titulo editable")
        self.assertTrue(title.font.bold)
        self.assertEqual(title.font.sz, 20)
        self.assertEqual(title.font.color.rgb[-6:], "FF0000")
        self.assertTrue(any(cell.value == "Detalle fuera de tabla" for cell in cells))
        formula = next(cell for cell in cells if str(cell.value).endswith("=1+1"))
        self.assertNotEqual(formula.data_type, "f")
        self.assertEqual(formula.font.color.rgb[-6:], "0000FF")
        header = next(cell for cell in cells if cell.value == "Cabecera combinada")
        self.assert_rgb(header.fill.fgColor.rgb[-6:], "CCE6FF")
        self.assertTrue(any(merged.min_row == header.row and merged.max_col == 2
                            for merged in sheet.merged_cells.ranges))
        self.assertEqual(len(workbook["Original 1"]._images), 1)
        with Image.open(io.BytesIO(workbook["Original 1"]._images[0]._data())) as reference:
            self.assertEqual(reference.size, (1000, 1300))
            self.assertEqual(reference.convert("RGB").getpixel((120, 600)), (0, 128, 0))

    def test_powerpoint_has_editable_text_table_and_background_without_old_text(self):
        output, _ = self.convert("pdf_to_pptx")
        presentation = Presentation(output)
        slide = presentation.slides[0]
        textboxes = [shape for shape in slide.shapes if shape.has_text_frame]
        title = next(shape for shape in textboxes if shape.text == "Titulo editable")
        run = title.text_frame.paragraphs[0].runs[0]
        self.assertTrue(run.font.bold)
        self.assertEqual(str(run.font.color.rgb), "FF0000")
        scale = presentation.slide_width / 500 / 12700
        self.assertAlmostEqual(run.font.size.pt, 20 * scale, delta=0.02)
        self.assertAlmostEqual(title.left / presentation.slide_width, 50 / 500, places=3)
        tables = [shape.table for shape in slide.shapes if shape.has_table]
        self.assertEqual(len(tables), 1)
        self.assertTrue(tables[0].cell(0, 0).is_merge_origin)
        self.assertEqual(tables[0].cell(0, 0).text, "Cabecera combinada")
        self.assert_rgb(str(tables[0].cell(0, 0).fill.fore_color.rgb), "CCE6FF")
        self.assertEqual(tables[0].cell(1, 1).text, "=1+1")
        with Image.open(io.BytesIO(slide.shapes[0].image.blob)) as background:
            # The title's old red pixels must not remain under the editable title.
            crop = background.convert("RGB").crop((90, 100, 400, 160))
            pixels = crop.tobytes()
            self.assertEqual(sum(pixels[i] > 150 and pixels[i + 1] < 100 and pixels[i + 2] < 100
                                 for i in range(0, len(pixels), 3)), 0)

    def test_scanned_and_rotated_pages_have_explicit_visual_fallback(self):
        for rotated in (False, True):
            with self.subTest(rotated=rotated):
                with pdf.open() as scanned:
                    page = scanned.new_page()
                    if rotated:
                        page.insert_text((50, 80), "TEXTO GIRADO")
                        page.set_rotation(90)
                    else:
                        image = Image.new("RGB", (100, 100), "blue")
                        buffer = io.BytesIO()
                        image.save(buffer, format="PNG")
                        page.insert_image(page.rect, stream=buffer.getvalue())
                        page.insert_text((50, 80), "OCR INVISIBLE", render_mode=3)
                    for operation in ("pdf_to_docx", "pdf_to_xlsx", "pdf_to_pptx"):
                        output, notice = export_office(operation, scanned, [0], self.folder, self.config)
                        self.assertGreater(output.stat().st_size, 0)
                        self.assertIn("1 paginas", notice)
                        self.assertIn("OCR", notice)

    def test_word_full_page_images_keep_one_editable_copy_of_tables(self):
        image = Image.new("RGB", (500, 650), "white")
        buffer = io.BytesIO()
        image.save(buffer, format="PNG")
        self.doc[0].insert_image(self.doc[0].rect, stream=buffer.getvalue(), overlay=False)
        output, notice = self.convert("pdf_to_docx")
        document = Document(output)
        self.assertEqual(len(document.tables), 1)
        text = "\n".join([paragraph.text for paragraph in document.paragraphs] +
                         [cell.text for cell in document.tables[0]._cells])
        self.assertEqual(text.count("Cabecera combinada"), 2)  # merged cell is reported once per grid column
        self.assertEqual(text.count("Titulo editable"), 1)
        self.assertNotIn("paginas complejas", notice)

    def test_word_keeps_real_document_pagination_and_word_positions(self):
        from collections import Counter
        from pdf_engine import convert_office, libreoffice_path

        if not libreoffice_path(self.config):
            self.skipTest("LibreOffice no disponible para validacion visual")
        with pdf.open() as source:
            for number in (1, 2):
                page = source.new_page(width=612, height=792)
                page.insert_text((270, 40), "PUBLICA", fontname="hebo", fontsize=11)
                page.insert_text((156, 62), "MINISTERIO DE PRUEBA NACIONAL", fontname="hebo", fontsize=9.96)
                page.insert_text((86, 120), "I.   ACTIVIDADES VERIFICADAS: texto con tabulacion", fontsize=12)
                page.insert_text((476, 50), " " * 40 + f"Pag. {number} de 2", fontsize=8)
                xs, ys = (92, 120, 204, 574), (300, 330, 420)
                for x in xs:
                    page.draw_line((x, ys[0]), (x, ys[-1]))
                for y in ys:
                    page.draw_line((xs[0], y), (xs[-1], y))
                page.insert_text((96, 320), "N", fontsize=9)
                page.insert_text((131, 320), "OPERACIONES", fontsize=9.96)
                page.insert_text((221, 320), "Encabezado de observacion", fontname="hebo", fontsize=9)
                page.insert_text((221, 350), "texto  justificado  con  espacios  amplios  hasta  el  borde", fontsize=9)
                page.insert_text((221, 362), "segunda linea normal de la celda", fontsize=9)
                page.insert_text((131, 350), "PROCESO DE", fontsize=9.96)
                watermark = page.new_shape()
                watermark.draw_polyline([(120, 700), (300, 260), (330, 270), (150, 710)])
                watermark.finish(color=None, fill=(0.85, 0.85, 0.85), fill_opacity=0.8)
                watermark.commit()
                signature = page.new_shape()
                signature.draw_bezier((240, 520), (260, 470), (300, 600), (340, 500))
                signature.finish(color=(0, 0, 0), width=2)
                signature.commit()
                page.insert_text((156, 740), "Carrera 1 No 2 - 3 Bogota", fontsize=7)
            work = self.folder / "pagination"
            work.mkdir()
            output, notice = export_office("pdf_to_docx", source, [0, 1], work, self.config)
            self.assertNotIn("paginas complejas", notice)
            self.assertEqual(len(Document(output).tables), 2)
            rendered = convert_office(output, work, self.config)
            with pdf.open(rendered) as result:
                self.assertEqual(len(result), 2)
                for original, converted in zip(source, result):
                    expected, actual = original.get_text("words"), converted.get_text("words")
                    self.assertEqual(Counter(word[4] for word in actual), Counter(word[4] for word in expected))
                    # Baselines are font independent; box tops vary with each font's ascender.
                    glyphs = [[(char["c"], char["bbox"][0], char["origin"][1])
                               for block in page.get_text("rawdict")["blocks"]
                               for line in block.get("lines", []) for span in line["spans"]
                               for char in span["chars"] if not char["c"].isspace()]
                              for page in (original, converted)]
                    for char, x, y in glyphs[0]:
                        error = min(max(abs(other[1] - x), abs(other[2] - y))
                                    for other in glyphs[1] if other[0] == char)
                        # Word only stores half-point sizes, so glyphs inside a word may drift slightly.
                        self.assertLess(error, 1.5, (char, x, y))
                    pixels = converted.get_pixmap(clip=(235, 470, 345, 600)).samples
                    self.assertGreater(sum(value < 60 for value in pixels[::3]), 50)

    def test_page_order_and_mixed_page_sizes(self):
        self.doc.new_page(width=650, height=500).insert_text((50, 80), "SEGUNDA")
        for operation in ("pdf_to_docx", "pdf_to_xlsx", "pdf_to_pptx"):
            with self.subTest(operation=operation):
                output, _ = export_office(operation, self.doc, [1, 0], self.folder, self.config)
                if operation == "pdf_to_docx":
                    result = Document(output)
                    self.assertEqual(len(result.sections), 2)
                    self.assertEqual(next(item.text for item in result.paragraphs if item.text), "SEGUNDA")
                elif operation == "pdf_to_xlsx":
                    workbook = load_workbook(output)
                    try:
                        self.assertEqual(workbook.sheetnames,
                                         ["Pagina 2", "Original 2", "Pagina 1", "Original 1"])
                    finally:
                        workbook.close()
                else:
                    presentation = Presentation(output)
                    self.assertEqual(len(presentation.slides), 2)
                    self.assertTrue(any(shape.has_text_frame and shape.text == "SEGUNDA"
                                        for shape in presentation.slides[0].shapes))

    def test_office_renderer_opens_outputs_without_losing_content(self):
        from pdf_engine import convert_office, libreoffice_path

        if not libreoffice_path(self.config):
            self.skipTest("LibreOffice no disponible para validacion visual")
        for operation in ("pdf_to_docx", "pdf_to_pptx"):
            with self.subTest(operation=operation):
                work = self.folder / operation
                work.mkdir()
                output, _ = export_office(operation, self.doc, [0], work, self.config)
                rendered = convert_office(output, work, self.config)
                with pdf.open(rendered) as result:
                    self.assertEqual(len(result), 1)
                    text = "".join(page.get_text() for page in result)
                    for expected in ("Titulo editable", "Cabecera combinada", "Detalle fuera de tabla", "=1+1"):
                        self.assertEqual(text.count(expected), 1, text)
                    pixmap = result[0].get_pixmap()
                    pixels = pixmap.samples
                    self.assertGreater(sum(pixels[i + 1] > 80 and pixels[i] < 50 and pixels[i + 2] < 50
                                           for i in range(0, len(pixels), 3)), 100)

    def test_http_conversion_keeps_editable_tables_and_fixed_export_names(self):
        from server import CLIENT_BUILD, create_app

        root = self.folder / "api"
        root.mkdir()
        app = create_app(root)
        self.addCleanup(app.extensions["instance_lock"].close)
        original = self.doc.tobytes()
        original_pixels = self.doc[0].get_pixmap().samples
        for operation, extension in (("pdf_to_docx", "docx"), ("pdf_to_xlsx", "xlsx"),
                                     ("pdf_to_pptx", "pptx")):
            with self.subTest(operation=operation):
                response = app.test_client().post(
                    "/api/process/" + operation,
                    headers={"Origin": "http://localhost", "X-PDF-Client": CLIENT_BUILD},
                    data={"options": "{}", "files": [(io.BytesIO(original), "formatted.pdf")]},
                )
                try:
                    self.assertEqual(response.status_code, 200, response.get_json() if response.is_json else "")
                    self.assertEqual(response.headers["Content-Disposition"], f"attachment; filename=EVA00.{extension}")
                    self.assertEqual(response.headers["X-PDF-Cleanup"], "complete")
                    self.assertIn("revise", response.headers["X-PDF-Notice"])
                    if extension == "docx":
                        self.assertEqual(len(Document(io.BytesIO(response.data)).tables), 1)
                    elif extension == "xlsx":
                        workbook = load_workbook(io.BytesIO(response.data))
                        try:
                            self.assertEqual(workbook.sheetnames, ["Pagina 1", "Original 1"])
                        finally:
                            workbook.close()
                    else:
                        slides = Presentation(io.BytesIO(response.data)).slides
                        self.assertEqual(sum(shape.has_table for shape in slides[0].shapes), 1)
                finally:
                    response.close()
                    response.request.environ["wsgi.input"].close()
        self.assertEqual(self.doc[0].get_pixmap().samples, original_pixels)
        self.assertEqual(list((root / "runtime").glob("job-*")), [])
