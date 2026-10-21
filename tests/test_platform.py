from __future__ import annotations

import io
import json
import sqlite3
import tempfile
import unittest
import zipfile
from concurrent.futures import ThreadPoolExecutor
from contextlib import closing
from pathlib import Path
from unittest.mock import patch

import pymupdf as pdf
from docx import Document
from openpyxl import Workbook, load_workbook
from PIL import Image
from pptx import Presentation

from isolation import InstanceLock
from pdf_engine import ProcessingError, capabilities, pages_from
from server import CLIENT_BUILD, ROOT, create_app


def sample_pdf(count=3, *, compressed=True):
    with pdf.open() as document:
        for index in range(count):
            page = document.new_page()
            page.insert_text((72, 90), f"DOCUMENTO DE PRUEBA {index + 1}", fontsize=22)
            page.insert_text((72, 140), "SECRETO 987654321", fontsize=16)
        return document.tobytes(deflate=compressed)


class PlatformTests(unittest.TestCase):
    @classmethod
    def setUpClass(cls):
        cls.temporary = tempfile.TemporaryDirectory(prefix="folio-test-")
        cls.root = Path(cls.temporary.name)
        cls.app = create_app(cls.root)
        cls.app.testing = True
        cls.client = cls.app.test_client()
        cls.config = cls.app.extensions["pdf_settings"]
        cls.headers = {"Origin": "http://localhost", "X-PDF-Client": CLIENT_BUILD}
        cls.original = sample_pdf()

    @classmethod
    def tearDownClass(cls):
        cls.app.extensions["instance_lock"].close()
        cls.temporary.cleanup()

    def tearDown(self):
        self.assertEqual(list((self.root / "runtime").glob("job-*")), [])

    def call(self, operation, options=None, uploads=None, expected=200):
        uploads = uploads if uploads is not None else [(self.original, "privado.pdf")]
        response = self.client.post(
            "/api/process/" + operation,
            headers=self.headers,
            data={"options": json.dumps(options or {}),
                  "files": [(io.BytesIO(data), name) for data, name in uploads]},
        )
        response.get_data()
        response.request.environ["wsgi.input"].close()
        self.assertEqual(response.status_code, expected, response.get_json() if response.is_json else response.data[:200])
        if expected == 200 and operation != "inspect":
            self.assertEqual(response.headers["X-PDF-Cleanup"], "complete")
            self.assertRegex(response.headers["Content-Disposition"],
                             r'filename="?EVA00\.[a-z0-9]+"?$')
            if response.mimetype == "application/zip":
                with zipfile.ZipFile(io.BytesIO(response.data)) as archive:
                    for name in archive.namelist():
                        self.assertRegex(Path(name).name, r"^EVA00_\d{2,}\.[a-z0-9]+$")
        self.assertEqual(response.headers["Cache-Control"], "no-store, private, max-age=0")
        return response

    def test_health_and_page(self):
        self.assertEqual(self.client.get("/health").json["client_build"], CLIENT_BUILD)
        response = self.client.get("/")
        self.assertIn(b"EVA PDF Enterprise", response.data)
        self.assertIn(b'id="multi-edit-tab"', response.data)
        self.assertIn("Edición múltiple".encode(), response.data)
        self.assertIn(b'id="session-tools"', response.data)
        self.assertIn("frame-ancestors 'none'", response.headers["Content-Security-Policy"])
        self.assertEqual(self.client.get("/api/capabilities").json["max_pages"], 0)

    def test_no_document_tables(self):
        with closing(sqlite3.connect(self.root / "instance" / "platform.sqlite3")) as connection:
            tables = connection.execute("SELECT name FROM sqlite_master WHERE type='table'").fetchall()
            self.assertEqual(tables, [("settings",)])
            values = str(connection.execute("SELECT * FROM settings").fetchall())
            self.assertNotIn("SECRETO", values)
            self.assertNotIn("privado.pdf", values)

    def test_desktop_section_under_construction(self):
        response = self.client.get("/")
        section = response.get_data(as_text=True).split('id="desktop-tools"', 1)[1].split("</section>", 1)[0]
        self.assertIn("Funcion en Construccion", section)
        self.assertIn('class="construction-actions" disabled', section)
        self.assertIn('id="camera-open" class="secondary" disabled', section)
        self.assertIn('id="screen-capture" class="secondary" disabled', section)
        self.assertNotIn('href="/desktop/download"', section)
        symbol = self.client.get("/static/construction.svg")
        self.assertEqual(symbol.status_code, 200)
        self.assertEqual(symbol.mimetype, "image/svg+xml")
        symbol.close()

    def test_origin_and_client_build(self):
        for headers in ({}, {"Origin": "http://externo", "X-PDF-Client": CLIENT_BUILD},
                        {"Origin": "http://localhost", "X-PDF-Client": "old"}):
            response = self.client.post("/api/process/merge", headers=headers)
            self.assertEqual(response.status_code, 403)

    def test_unknown_operation(self):
        self.call("nonexistent", expected=404)

    def test_merge(self):
        response = self.call("merge", uploads=[(self.original, "one.pdf"), (sample_pdf(1), "two.pdf")])
        with pdf.open(stream=response.data, filetype="pdf") as document:
            self.assertEqual(len(document), 4)

    def test_split(self):
        response = self.call("split", {"groups": "3;1-2"})
        with zipfile.ZipFile(io.BytesIO(response.data)) as archive:
            self.assertEqual(len(archive.namelist()), 2)
            with pdf.open(stream=archive.read(archive.namelist()[0]), filetype="pdf") as doc:
                self.assertEqual(len(doc), 1)
                self.assertIn("PRUEBA 3", doc[0].get_text())
            with pdf.open(stream=archive.read(archive.namelist()[1]), filetype="pdf") as doc:
                self.assertEqual(len(doc), 2)

    def test_organize(self):
        response = self.call("organize", {"pages": "3,1"})
        with pdf.open(stream=response.data, filetype="pdf") as document:
            self.assertEqual(len(document), 2)
            self.assertIn("PRUEBA 3", document[0].get_text())
            self.assertIn("PRUEBA 1", document[1].get_text())

    def test_rotate(self):
        response = self.call("rotate", {"pages": "2", "angle": "90"})
        with pdf.open(stream=response.data, filetype="pdf") as document:
            self.assertEqual([page.rotation for page in document], [0, 90, 0])

    def test_bad_pages(self):
        self.call("organize", {"pages": "0,50"}, expected=422)
        self.call("rotate", {"pages": "1,1"}, expected=422)
        self.call("rotate", {"angle": "91"}, expected=422)
        self.call("split", {"groups": "1;;2"}, expected=422)

    def test_page_parser(self):
        self.assertEqual(pages_from("3-1", 3), [2, 1, 0])
        for value in ("1,,2", "4", "1,1", "1;a"):
            with self.assertRaises(ProcessingError):
                pages_from(value, 3)

    def test_compress_preserves_text(self):
        response = self.call("compress")
        with pdf.open(stream=response.data, filetype="pdf") as document:
            self.assertIn("SECRETO", document[0].get_text())
            self.assertEqual(len(document), 3)
        self.assertLessEqual(len(response.data), len(self.original))

    def test_compress_measured_reduction(self):
        with pdf.open() as doc:
            image = Image.effect_noise((1200, 1200), 40).convert("RGB")
            buf = io.BytesIO()
            image.save(buf, format="BMP")
            page = doc.new_page()
            page.insert_image(page.rect, stream=buf.getvalue())
            source = doc.tobytes(deflate=False)
        response = self.call("compress", {"compression": "raster"}, [(source, "large.pdf")])
        self.assertLess(len(response.data), len(source) * 0.5)
        self.assertIn("Compresion visual", response.headers["X-PDF-Notice"])

    def test_images_to_pdf(self):
        image = Image.new("RGBA", (200, 100), (255, 0, 0, 120))
        data = io.BytesIO()
        image.save(data, format="PNG")
        response = self.call("images_to_pdf", uploads=[(data.getvalue(), "a.png"), (data.getvalue(), "b.png")])
        with pdf.open(stream=response.data, filetype="pdf") as doc:
            self.assertEqual(len(doc), 2)
            self.assertEqual(doc[0].rect.width, 150)

    def test_pdf_to_images(self):
        response = self.call("pdf_to_images", {"pages": "2", "dpi": "72"})
        with zipfile.ZipFile(io.BytesIO(response.data)) as archive:
            self.assertEqual(len(archive.namelist()), 1)
            self.assertTrue(archive.namelist()[0].endswith(".png"))
            with Image.open(io.BytesIO(archive.read(archive.namelist()[0]))) as image:
                self.assertEqual(image.width, 595)

    def test_extract_text(self):
        response = self.call("extract_text", {"pages": "2"})
        text = response.data.decode()
        self.assertIn("PRUEBA 2", text)
        self.assertNotIn("PRUEBA 1", text)

    def test_preview(self):
        response = self.call("inspect", {"page": "2"})
        self.assertEqual(response.json["pages"], 3)
        self.assertEqual(response.json["page"], 2)
        self.assertGreater(len(response.json["image"]), 100)

    def test_editor_text_and_real_redaction(self):
        # El secreto ocupa y=123..145; se elimina completamente esa zona.
        response = self.call("edit", {"edits": [
            {"kind": "redact", "page": 1, "x": 0.10, "y": 0.13, "w": 0.65, "h": 0.08},
            {"kind": "text", "page": 2, "x": 0.1, "y": 0.3, "w": 0.8, "h": 0.15, "text": "NUEVO TEXTO", "size": 16},
        ]})
        with pdf.open(stream=response.data, filetype="pdf") as doc:
            self.assertNotIn("SECRETO", doc[0].get_text())
            self.assertNotIn("987654321", doc[0].get_text())
            self.assertIn("SECRETO", doc[1].get_text())
            self.assertIn("NUEVO TEXTO", doc[1].get_text())
            pix = doc[0].get_pixmap()
            self.assertEqual(pix.pixel(100, 130), (0, 0, 0))

    def test_invalid_edits(self):
        self.call("edit", {"edits": []}, expected=422)
        self.call("edit", {"edits": [
            {"kind": "redact", "page": 1, "x": 0.8, "y": 0.1, "w": 0.8, "h": 0.1},
        ]}, expected=422)

    def test_watermark(self):
        response = self.call("watermark", {"pages": "1", "text": "COPIA"})
        with pdf.open(stream=response.data, filetype="pdf") as document:
            self.assertIn("COPIA", document[0].get_text())
            self.assertNotIn("COPIA", document[1].get_text())

    def test_protect_and_unlock(self):
        response = self.call("protect", {"new_password": "clave-segura-123"})
        with pdf.open(stream=response.data, filetype="pdf") as doc:
            self.assertTrue(doc.needs_pass)
            self.assertFalse(doc.authenticate("incorrecta"))
            self.assertTrue(doc.authenticate("clave-segura-123"))
        self.call("unlock", {"password": "incorrecta"}, [(response.data, "encrypted.pdf")], expected=422)
        unlocked = self.call("unlock", {"password": "clave-segura-123"}, [(response.data, "encrypted.pdf")])
        with pdf.open(stream=unlocked.data, filetype="pdf") as doc:
            self.assertFalse(doc.needs_pass)
            self.assertIn("SECRETO", doc[0].get_text())
        self.call("protect", {"new_password": "123"}, expected=422)

    def test_pdf_to_word(self):
        response = self.call("pdf_to_docx", {"pages": "1"})
        document = Document(io.BytesIO(response.data))
        self.assertIn("SECRETO", "\n".join(p.text for p in document.paragraphs))

    def test_pdf_to_excel(self):
        response = self.call("pdf_to_xlsx", {"pages": "1"})
        workbook = load_workbook(io.BytesIO(response.data))
        try:
            self.assertEqual(workbook.sheetnames, ["Pagina 1", "Original 1"])
            self.assertTrue(any("SECRETO" in str(cell.value) for row in workbook.active for cell in row))
        finally:
            workbook.close()

    def test_pdf_to_powerpoint(self):
        response = self.call("pdf_to_pptx", {"pages": "2,1"})
        presentation = Presentation(io.BytesIO(response.data))
        self.assertEqual(len(presentation.slides), 2)
        self.assertGreater(len(presentation.slides[0].shapes), 1)
        self.assertTrue(any(shape.has_text_frame and "PRUEBA 2" in shape.text
                            for shape in presentation.slides[0].shapes))

    def test_office_formats_to_pdf(self):
        if not capabilities(self.config, ROOT)["office"]:
            self.skipTest("LibreOffice no disponible")
        word = Document()
        word.add_paragraph("PRUEBA WORD LOCAL")
        word_data = io.BytesIO()
        word.save(word_data)
        excel = Workbook()
        excel.active["A1"] = "PRUEBA EXCEL LOCAL"
        excel_data = io.BytesIO()
        excel.save(excel_data)
        slides = Presentation()
        slide = slides.slides.add_slide(slides.slide_layouts[0])
        slide.shapes.title.text = "PRUEBA PPT LOCAL"
        slides_data = io.BytesIO()
        slides.save(slides_data)
        for stream, name, expected in (
            (word_data, "word.docx", "WORD"), (excel_data, "book.xlsx", "EXCEL"),
            (slides_data, "slides.pptx", "PPT"),
        ):
            with self.subTest(name=name):
                response = self.call("office_to_pdf", uploads=[(stream.getvalue(), name)])
                with pdf.open(stream=response.data, filetype="pdf") as doc:
                    self.assertIn(expected, "".join(page.get_text() for page in doc))

    def test_ocr_local(self):
        if not capabilities(self.config, ROOT)["ocr"]:
            self.skipTest("Modelos OCR no disponibles")
        with pdf.open(stream=sample_pdf(1), filetype="pdf") as original:
            image = original[0].get_pixmap(dpi=180).tobytes("png")
        with pdf.open() as scanned:
            page = scanned.new_page()
            page.insert_image(page.rect, stream=image)
            source = scanned.tobytes()
        self.call("extract_text", uploads=[(source, "scan.pdf")], expected=422)
        response = self.call("ocr", {"dpi": "180"}, [(source, "scan.pdf")])
        with pdf.open(stream=response.data, filetype="pdf") as doc:
            text = doc[0].get_text().upper()
            self.assertIn("DOCUMENTO", text)
            self.assertIn("987654321", text)
        self.assertIn("OCR automatico", response.headers["X-PDF-Notice"])

    def test_bad_document_and_payload(self):
        self.call("rotate", uploads=[(b"not a PDF", "bad.pdf")], expected=422)
        self.call("rotate", uploads=[(b"text", "file.exe")], expected=422)
        self.call("rotate", options=["bad"], expected=400)
        self.call("rotate", uploads=[], expected=400)

    def test_page_and_output_limits(self):
        original_limit = self.config["max_pages"]
        self.config["max_pages"] = 1
        try:
            self.call("rotate", expected=422)
        finally:
            self.config["max_pages"] = original_limit
        original_limit = self.config["max_output_mb"]
        self.config["max_output_mb"] = 0.00001
        try:
            self.call("rotate", expected=413)
        finally:
            self.config["max_output_mb"] = original_limit

    def test_busy_does_not_stage_files(self):
        gate = self.app.extensions["job_gate"]
        for _ in range(self.config["max_parallel_jobs"]):
            gate.acquire()
        try:
            with patch.dict(self.config, queue_wait_seconds=.02):
                self.call("rotate", expected=503)
            self.assertEqual(list((self.root / "runtime").glob("job-*")), [])
        finally:
            for _ in range(self.config["max_parallel_jobs"]):
                gate.release()

    def test_parallel_jobs_are_isolated(self):
        def execute(number):
            with self.app.test_client() as client:
                response = client.post("/api/process/extract_text", headers=self.headers, data={
                    "options": json.dumps({"pages": str(number)}),
                    "files": (io.BytesIO(self.original), "same-name.pdf"),
                })
                data = response.get_data().decode()
                response.request.environ["wsgi.input"].close()
                return response.status_code, data
        with ThreadPoolExecutor(max_workers=2) as executor:
            results = list(executor.map(execute, (1, 3)))
        self.assertEqual([result[0] for result in results], [200, 200])
        self.assertIn("PRUEBA 1", results[0][1])
        self.assertNotIn("PRUEBA 3", results[0][1])
        self.assertIn("PRUEBA 3", results[1][1])

    def test_cleanup_happens_before_response_creation(self):
        from flask import send_file
        def checked_send_file(*args, **kwargs):
            self.assertEqual(list((self.root / "runtime").glob("job-*")), [])
            return send_file(*args, **kwargs)
        with patch("server.send_file", side_effect=checked_send_file):
            self.call("rotate")

    def test_large_multipart_is_spooled_to_runtime(self):
        from werkzeug.test import EnvironBuilder
        payload = b"private" * 200000
        builder = EnvironBuilder(method="POST", data={"files": (io.BytesIO(payload), "large.pdf")})
        environment = builder.get_environ()
        incoming = self.app.request_class(environment)
        spool = self.root / "runtime" / "spool"
        try:
            stream = incoming.files["files"].stream
            self.assertNotIsInstance(stream, io.BytesIO)
            self.assertEqual(Path(stream.name).parent, spool)
            self.assertEqual(incoming.files["files"].read(), payload)
        finally:
            incoming.close()
            environment["wsgi.input"].close()
            builder.close()
        self.assertEqual(list(spool.iterdir()), [])

    def test_per_file_limit(self):
        self.assertEqual(self.client.get("/api/capabilities").json["max_file_mb"], 1024)
        original_limit = self.config["max_file_mb"]
        self.config["max_file_mb"] = 0.001
        try:
            self.call("rotate", expected=413)
        finally:
            self.config["max_file_mb"] = original_limit
        self.call("rotate")

    def test_session_quota_five_files_one_gb(self):
        caps = self.client.get("/api/capabilities").json
        self.assertEqual((caps["max_files"], caps["max_upload_mb"], caps["max_file_mb"]), (5, 1024, 1024))
        many = [(self.original, f"parte-{index}.pdf") for index in range(6)]
        self.call("watermark", options={"text": "X"}, uploads=many[:2], expected=422)
        self.call("overlay", uploads=many, expected=400)
        original_total = self.config["max_upload_mb"]
        # Cada archivo cabe solo, pero el conjunto excede la cuota de la sesion.
        self.config["max_upload_mb"] = len(self.original) * 1.5 / (1024 * 1024)
        try:
            self.call("rotate")
            self.call("merge", uploads=many[:2], expected=413)
        finally:
            self.config["max_upload_mb"] = original_total

    def test_bulk_tools_accept_many_files_within_quota(self):
        caps = self.client.get("/api/capabilities").json
        self.assertIn("compress", caps["bulk_operations"])
        self.assertEqual(caps["max_batch_files"], 200)
        many = [(sample_pdf(index % 3 + 1), f"doc-{index}.pdf") for index in range(12)]
        merged = self.call("merge", uploads=many)
        with pdf.open(stream=merged.data, filetype="pdf") as document:
            self.assertEqual(len(document), sum(index % 3 + 1 for index in range(12)))
        compressed = self.call("compress", uploads=many)
        with zipfile.ZipFile(io.BytesIO(compressed.data)) as archive:
            names = archive.namelist()
            self.assertEqual(len(names), 12)
            self.assertTrue(all(name.endswith(".pdf") for name in names))
            with pdf.open(stream=archive.read(names[2]), filetype="pdf") as doc:
                self.assertEqual(len(doc), 3)
        self.assertIn("12 archivos", compressed.headers["X-PDF-Notice"])
        split = self.call("split", uploads=many[:3])
        with zipfile.ZipFile(io.BytesIO(split.data)) as archive:
            self.assertEqual([Path(name).parent.name for name in archive.namelist()],
                             ["01", "02", "02", "03", "03", "03"])
        text = self.call("extract_text", uploads=many[:2])
        with zipfile.ZipFile(io.BytesIO(text.data)) as archive:
            self.assertIn("SECRETO", archive.read(archive.namelist()[1]).decode())
        for operation in ("pdf_to_docx", "compress", "split"):
            with self.subTest(operation=operation):
                broken = self.call(operation, uploads=[many[0], (b"not a PDF", "bad.pdf")], expected=422)
                self.assertIn("Archivo 2", broken.json["error"])
                self.assertNotIn("bad.pdf", broken.get_data(as_text=True))
        original_limit = self.config["max_batch_files"]
        self.config["max_batch_files"] = 3
        try:
            self.call("compress", uploads=many[:4], expected=400)
        finally:
            self.config["max_batch_files"] = original_limit

    def test_large_result_is_streamed_and_deleted(self):
        outgoing = self.root / "runtime" / "outgoing"
        with patch("server.INLINE_RESULT_BYTES", 1):
            response = self.client.post(
                "/api/process/rotate", headers=self.headers,
                data={"options": "{}", "files": [(io.BytesIO(self.original), "privado.pdf")]},
            )
            self.assertEqual(response.status_code, 200)
            self.assertEqual(response.headers["X-PDF-Cleanup"], "on-delivery")
            self.assertEqual(list((self.root / "runtime").glob("job-*")), [])
            self.assertEqual(len(list(outgoing.iterdir())), 1)
            data = response.get_data()
            self.assertEqual(int(response.headers["Content-Length"]), len(data))
            response.close()
            response.request.environ["wsgi.input"].close()
        with pdf.open(stream=data, filetype="pdf") as document:
            self.assertEqual(len(document), 3)
        self.assertEqual(list(outgoing.iterdir()), [])

    def test_timeout_cleans_job(self):
        original_limit = self.config["timeout_seconds"]
        self.config["timeout_seconds"] = 0.001
        try:
            self.call("rotate", expected=408)
        finally:
            self.config["timeout_seconds"] = original_limit
        self.call("rotate")

    def test_oversize_body_does_not_exhaust_gate(self):
        previous = self.app.config["MAX_CONTENT_LENGTH"]
        self.app.config["MAX_CONTENT_LENGTH"] = 512
        try:
            for _ in range(3):
                self.call("rotate", expected=413)
        finally:
            self.app.config["MAX_CONTENT_LENGTH"] = previous
        self.call("rotate")

    def test_cleanup_failure_blocks_and_reports(self):
        real_rmtree = __import__("shutil").rmtree
        try:
            with patch("server.shutil.rmtree", side_effect=PermissionError("blocked")):
                self.call("rotate", expected=503)
            self.assertEqual(self.client.get("/health").status_code, 503)
            self.call("rotate", expected=503)
        finally:
            self.app.extensions.pop("cleanup_failed", None)
            self.app.extensions["workload_queue"].closed = False
            for folder in (self.root / "runtime").glob("job-*"):
                real_rmtree(folder)

    def test_instance_lock(self):
        with self.assertRaises(RuntimeError):
            InstanceLock(self.root / "runtime" / "instance.lock")

    def test_startup_removes_abandoned_job(self):
        with tempfile.TemporaryDirectory(prefix="folio-restart-") as name:
            root = Path(name)
            abandoned = root / "runtime" / "job-abandoned"
            abandoned.mkdir(parents=True)
            (abandoned / "secret.pdf").write_bytes(self.original)
            app = create_app(root)
            try:
                self.assertFalse(abandoned.exists())
            finally:
                app.extensions["instance_lock"].close()


if __name__ == "__main__":
    unittest.main()
