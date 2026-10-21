from __future__ import annotations

import io
from pathlib import Path

import pymupdf as pdf


def contains(rect, bbox) -> bool:
    box = pdf.Rect(bbox)
    return pdf.Rect(rect).contains((box.tl + box.br) / 2)


def straight_path(drawing: dict) -> bool:
    for item in drawing["items"]:
        if item[0] == "c" or (item[0] == "qu" and not item[1].is_rectangular):
            return False
        if item[0] == "l" and abs(item[1].x - item[2].x) > 0.5 and abs(item[1].y - item[2].y) > 0.5:
            return False
    return True


def font_name(span: dict) -> str:
    name = span.get("font", "Arial").split("+")[-1]
    for prefix, family in (("Arial", "Arial"), ("TimesNewRoman", "Times New Roman"),
                           ("CourierNew", "Courier New"), ("Calibri", "Calibri"),
                           ("Cambria", "Cambria")):
        if name.startswith(prefix):
            return family
    return name.replace("-BoldItalic", "").replace("-Bold", "").replace("-Italic", "")


def rgb(span: dict) -> str:
    return f'{span.get("color", 0) & 0xFFFFFF:06X}'


def lines_in(lines: list[dict], rect) -> list[dict]:
    return [
        {**line, "spans": [span for span in line["spans"] if contains(rect, span["bbox"])]}
        for line in lines if any(contains(rect, span["bbox"]) for span in line["spans"])
    ]


def page_layout(page: pdf.Page, config: dict) -> dict:
    from pdf_engine import render

    visible = [trace["bbox"] for trace in page.get_texttrace()
               if trace["type"] != 3 and trace.get("opacity", 1) > 0]
    if page.rotation or not visible:
        return {"fallback": True, "image": render(page, 144, config).tobytes("png")}
    blocks = page.get_text("rawdict", flags=pdf.TEXTFLAGS_DICT & ~pdf.TEXT_PRESERVE_IMAGES)["blocks"]
    lines = []
    for block in blocks:
        for line in block.get("lines", []):
            for span in line["spans"]:
                span["text"] = "".join(char["c"] for char in span.get("chars", []))
            spans = [span for span in line["spans"] if span["text"].strip()
                     and span.get("alpha", 255) > 0
                     and any(contains(rect, span["bbox"]) for rect in visible)]
            if spans:
                lines.append({**line, "spans": spans})
    # Table objects must be consumed before loading another page.
    # Table rules are straight; curves and diagonals (watermarks, signatures) must not split grids.
    finder = page.find_tables(paths=[drawing for drawing in page.get_drawings() if straight_path(drawing)])
    tables = []
    for table in finder.tables:
        cells = [rect for row in table.rows for rect in row.cells if rect]
        xs = sorted({round(rect[i], 2) for rect in cells for i in (0, 2)})
        ys = sorted({round(rect[i], 2) for rect in cells for i in (1, 3)})
        if len(xs) < 2 or len(ys) < 2:
            continue
        grid = []
        for rect in cells:
            x0, y0, x1, y1 = (round(value, 2) for value in rect)
            grid.append((ys.index(y0), xs.index(x0), ys.index(y1) - 1, xs.index(x1) - 1, rect))
        tables.append({"bbox": table.bbox, "xs": xs, "ys": ys, "cells": grid})
    visuals = []
    for info in page.get_image_info():
        rect = pdf.Rect(info["bbox"]) & page.rect
        if not rect.is_empty:
            visuals.append(tuple(rect))
    for rect in page.cluster_drawings():
        if not any(pdf.Rect(table["bbox"]).contains(rect) for table in tables):
            if rect.width > 2 and rect.height > 2 and not any(pdf.Rect(box).contains(rect) for box in visuals):
                visuals.append(tuple(rect & page.rect))
    for line in lines:
        if tuple(line["dir"]) != (1.0, 0.0) and not any(
            pdf.Rect(box).contains(pdf.Rect(line["bbox"])) for box in visuals
        ):
            visuals.append(line["bbox"])
    return {"fallback": False, "lines": lines, "tables": tables, "visuals": visuals}


def crop_png(page: pdf.Page, rect, config: dict) -> bytes:
    from pdf_engine import ProcessingError

    box = pdf.Rect(rect)
    if box.width * box.height * 4 > config["max_render_pixels"]:
        raise ProcessingError("Grafico demasiado grande para exportar. Procese una pagina mas pequena.")
    return page.get_pixmap(clip=box, dpi=144, colorspace=pdf.csRGB, alpha=False).tobytes("png")


def cell_fill(drawings: list[dict], rect) -> str:
    center = (pdf.Rect(rect).tl + pdf.Rect(rect).br) / 2
    color = "FFFFFF"
    for drawing in drawings:
        fill = drawing.get("fill")
        if fill and len(fill) == 3 and drawing.get("fill_opacity", 1) == 1 and drawing["rect"].contains(center):
            color = "".join(f"{max(0, min(255, round(channel * 255))):02X}" for channel in fill)
    return color


def metric_font(span: dict):
    """Return a built-in font with metrics compatible with the PDF font, when known."""
    family = font_name(span)
    base = {"Arial": "he", "Times New Roman": "ti", "Courier New": "co"}.get(family)
    base = base or next((code for prefix, code in (("Helvetica", "he"), ("Times", "ti"), ("Courier", "co"))
                         if family.startswith(prefix)), None)
    if not base:
        return None
    fonts = {"he": ("helv", "hebo", "heit", "hebi"), "ti": ("tiro", "tibo", "tiit", "tibi"),
             "co": ("cour", "cobo", "coit", "cobi")}
    return pdf.Font(fonts[base][bool(span["flags"] & 16) + 2 * bool(span["flags"] & 2)])


def word_tracking(run, span: dict, text: str, size: float, target: float) -> float:
    """Match a PDF word or gap width despite Word's half-point sizes; return the width Word will use."""
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn

    font = metric_font(span)
    if not font or not text:
        return target
    natural = font.text_length(text, fontsize=size)
    # Word stores character spacing in whole twips and caps it at 1584 twips per character.
    limit = 79.2 if text.isspace() else size * 0.5
    twips = round(max(-size * 0.2, min(limit, (target - natural) / len(text))) * 20)
    if twips:
        spacing = OxmlElement("w:spacing")
        spacing.set(qn("w:val"), str(twips))
        run._r.get_or_add_rPr().append(spacing)
    return natural + twips * len(text) / 20


def word_line_box(spans: list[dict]) -> tuple[float, float, float, float]:
    """Return the line box Word uses: baseline from the PDF, ascent/descent from the substituted font."""
    metrics = {"he": (0.905, 0.212), "ti": (0.891, 0.216), "co": (0.833, 0.300)}
    x0 = min(span["bbox"][0] for span in spans)
    x1 = max(span["bbox"][2] for span in spans)
    baseline = max(span["origin"][1] for span in spans) if all("origin" in span for span in spans) else None
    if baseline is None:
        return x0, min(span["bbox"][1] for span in spans), x1, max(span["bbox"][3] for span in spans)
    ascent = descent = 0.0
    for span in spans:
        font = metric_font(span)
        size = max(0.5, round(span["size"] * 2) / 2)
        up, down = metrics[font.name[:2].lower()] if font and font.name[:2].lower() in metrics else (
            (baseline - span["bbox"][1]) / size, (span["bbox"][3] - baseline) / size)
        ascent, descent = max(ascent, up * size), max(descent, down * size)
    return x0, baseline - ascent, x1, baseline + descent


def word_segments(span: dict) -> list[tuple[str, float, float, bool]]:
    """Split a span into words with exact PDF positions and whether whitespace precedes each word."""
    from pdf_engine import clean_text

    if not span.get("chars"):
        return [(clean_text(span["text"]).strip(), span["bbox"][0], span["bbox"][2], span["text"][:1].isspace())]
    segments, word, start, end, spaced = [], "", 0.0, 0.0, False
    for char in span["chars"]:
        if char["c"].isspace():
            if word:
                segments.append((clean_text(word), start, end, spaced))
                word = ""
            spaced = True
            continue
        if not word:
            start = char["bbox"][0]
        word += char["c"]
        end = char["bbox"][2]
    if word:
        segments.append((clean_text(word), start, end, spaced))
    return [segment for i, segment in enumerate(segments) if segment[0] or i == 0]


def docx_runs(paragraph, lines: list[dict]) -> None:
    import math
    from docx.shared import Pt, RGBColor

    def styled(text: str, span: dict, target: float) -> float:
        if text == " ":
            text = " " * max(1, math.ceil(target / 70))
        run = paragraph.add_run(text)
        size = max(0.5, round(span["size"] * 2) / 2)
        font = metric_font(span)
        if text.isspace() and font:
            # Renderers may ignore negative spacing on spaces; a smaller space keeps tight gaps exact.
            unit = font.text_length(" ", fontsize=1)
            size = min(size, max(1.0, math.floor(target / len(text) / unit * 2) / 2))
        run.font.name = font_name(span)
        run.font.size = Pt(size)
        run.font.bold = bool(span["flags"] & 16)
        run.font.italic = bool(span["flags"] & 2)
        run.font.superscript = bool(span["flags"] & 1)
        run.font.color.rgb = RGBColor.from_string(rgb(span))
        return word_tracking(run, span, text, size, target)

    for index, line in enumerate(lines):
        if index:
            paragraph.add_run().add_break()
        cursor = previous = None
        # Gaps are measured from Word's emitted position so rounding never accumulates.
        position = min(span["bbox"][0] for span in line["spans"])
        for span in line["spans"]:
            for order, (text, x0, x1, spaced) in enumerate(word_segments(span)):
                if not text:
                    continue
                spaced = spaced or (not order and cursor is not None and x0 - cursor > 1)
                if (cursor is None and x0 - position > 0.5) or (cursor is not None and spaced):
                    position += styled(" ", previous or span, max(0.0, x0 - position))
                position += styled(text, span, x1 - x0)
                cursor, previous = x1, span


def word_background(paragraph, image: bytes, width: float, height: float) -> None:
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn
    from docx.shared import Pt

    inline = paragraph.add_run().add_picture(io.BytesIO(image), width=Pt(width), height=Pt(height))._inline
    anchor = OxmlElement("wp:anchor")
    for key, value in {"distT": "0", "distB": "0", "distL": "0", "distR": "0",
                       "simplePos": "0", "relativeHeight": "0", "behindDoc": "1",
                       "locked": "0", "layoutInCell": "1", "allowOverlap": "1"}.items():
        anchor.set(key, value)
    simple = OxmlElement("wp:simplePos")
    simple.set("x", "0")
    simple.set("y", "0")
    anchor.append(simple)
    for axis in ("H", "V"):
        position = OxmlElement(f"wp:position{axis}")
        position.set("relativeFrom", "page")
        offset = OxmlElement("wp:posOffset")
        offset.text = "0"
        position.append(offset)
        anchor.append(position)
    anchor.append(inline.find(qn("wp:extent")))
    anchor.append(OxmlElement("wp:wrapNone"))
    for tag in ("wp:docPr", "wp:cNvGraphicFramePr", "a:graphic"):
        anchor.append(inline.find(qn(tag)))
    inline.getparent().replace(inline, anchor)


def word_frame(paragraph, bbox, page_width: float) -> None:
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn
    from docx.shared import Pt

    x0, y0, x1, y1 = bbox
    paragraph.paragraph_format.space_before = Pt(0)
    paragraph.paragraph_format.space_after = Pt(0)
    paragraph.paragraph_format.line_spacing = Pt(max(1, y1 - y0))
    frame = OxmlElement("w:framePr")
    for key, value in {
        "x": str(round(x0 * 20)), "y": str(round(y0 * 20)),
        "w": str(round(max(1, page_width - x0) * 20)),
        "hAnchor": "page", "vAnchor": "page", "wrap": "none",
    }.items():
        frame.set(qn("w:" + key), value)
    paragraph._p.get_or_add_pPr().append(frame)


def word_table(result, content: dict, lines: list[dict]) -> None:
    from docx.enum.table import WD_ROW_HEIGHT_RULE
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn
    from docx.shared import Pt

    xs, ys = content["xs"], content["ys"]
    table = result.add_table(rows=len(ys) - 1, cols=len(xs) - 1)
    table.autofit = False
    properties = table._tbl.tblPr
    position = OxmlElement("w:tblpPr")
    for key, value in {
        "horzAnchor": "page", "vertAnchor": "page",
        "tblpX": str(round(xs[0] * 20)), "tblpY": str(round(ys[0] * 20)),
        "leftFromText": "0", "rightFromText": "0", "topFromText": "0", "bottomFromText": "0",
    }.items():
        position.set(qn("w:" + key), value)
    properties.append(position)
    margins = OxmlElement("w:tblCellMar")
    borders = OxmlElement("w:tblBorders")
    for edge in ("top", "left", "bottom", "right"):
        margin = OxmlElement("w:" + edge)
        margin.set(qn("w:w"), "0")
        margin.set(qn("w:type"), "dxa")
        margins.append(margin)
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        border = OxmlElement("w:" + edge)
        border.set(qn("w:val"), "nil")
        borders.append(border)
    properties.append(margins)
    properties.append(borders)
    for column, width in zip(table.columns, [b - a for a, b in zip(xs, xs[1:])]):
        column.width = Pt(width)
    for row, start, end_row, end_col, rect in content["cells"]:
        cell = table.cell(row, start)
        if end_row != row or end_col != start:
            cell = cell.merge(table.cell(end_row, end_col))
        cell.width = Pt(rect[2] - rect[0])
        cell_lines = sorted(lines_in(lines, rect), key=lambda line: (line["bbox"][1], line["bbox"][0]))
        bottom = rect[1]
        for index, line in enumerate(cell_lines):
            paragraph = cell.paragraphs[0] if not index else cell.add_paragraph()
            box = pdf.Rect(word_line_box(line["spans"]))
            paragraph.paragraph_format.space_before = Pt(max(0, box.y0 - bottom))
            paragraph.paragraph_format.space_after = Pt(0)
            text_width = pdf.Rect(line["spans"][0]["bbox"]).x0
            text_width = max(pdf.Rect(span["bbox"]).x1 - (len(span["text"]) - len(span["text"].rstrip())) * span["size"] * 0.28
                             for span in line["spans"]) - text_width
            paragraph.paragraph_format.left_indent = Pt(max(0, min(box.x0 - rect[0], rect[2] - rect[0] - text_width - 1)))
            paragraph.paragraph_format.line_spacing = Pt(max(1, box.height))
            docx_runs(paragraph, [line])
            bottom = box.y1
        if not cell_lines:
            cell.paragraphs[0].paragraph_format.space_after = Pt(0)
            cell.paragraphs[0].paragraph_format.line_spacing = Pt(1)
    for row, height in zip(table.rows, [b - a for a, b in zip(ys, ys[1:])]):
        row.height = Pt(height)
        row.height_rule = WD_ROW_HEIGHT_RULE.EXACTLY


def word_document(doc: pdf.Document, selected: list[int], output: Path, config: dict) -> int:
    from docx import Document
    from docx.enum.section import WD_SECTION_START
    from docx.shared import Pt
    from pdf_engine import render

    result = Document()
    fallback = 0
    for index, number in enumerate(selected):
        page = doc[number]
        section = result.sections[0] if not index else result.add_section(WD_SECTION_START.NEW_PAGE)
        section.page_width, section.page_height = Pt(page.rect.width), Pt(page.rect.height)
        section.left_margin = section.right_margin = Pt(0)
        section.top_margin = section.bottom_margin = Pt(0)
        anchor = result.add_paragraph()
        anchor.paragraph_format.space_before = anchor.paragraph_format.space_after = Pt(0)
        anchor.paragraph_format.line_spacing = Pt(1)
        layout = page_layout(page, config)
        if layout["fallback"]:
            word_background(anchor, layout["image"], page.rect.width, page.rect.height)
            fallback += 1
            continue
        tables = [table for table in layout["tables"] if len(table["xs"]) <= 64 and not any(
            tuple(line["dir"]) != (1.0, 0.0) and contains(table["bbox"], line["bbox"])
            for line in layout["lines"]
        )]
        lines = [line for line in layout["lines"] if tuple(line["dir"]) == (1.0, 0.0)]
        # Retain marks, signatures and table rules, removing only text rebuilt as editable objects.
        with pdf.open() as background:
            background.insert_pdf(doc, from_page=number, to_page=number)
            backdrop = background[0]
            for line in lines:
                for span in line["spans"]:
                    backdrop.add_redact_annot(span["bbox"], fill=False)
            backdrop.apply_redactions(images=0, graphics=0)
            image = render(backdrop, 144, config).tobytes("png")
        word_background(anchor, image, page.rect.width, page.rect.height)
        for line in sorted(lines, key=lambda item: (item["bbox"][1], item["bbox"][0])):
            spans = [span for span in line["spans"]
                     if not any(contains(table["bbox"], span["bbox"]) for table in tables)]
            if not spans:
                continue
            box = pdf.Rect(word_line_box(spans))
            paragraph = result.add_paragraph()
            word_frame(paragraph, tuple(box), page.rect.width)
            docx_runs(paragraph, [{**line, "spans": spans}])
        for table in tables:
            word_table(result, table, lines)
        # A normal paragraph keeps floating tables anchored to their source section.
        end = result.add_paragraph()
        end.paragraph_format.space_after = Pt(0)
        end.paragraph_format.line_spacing = Pt(1)
    result.save(output)
    return fallback


def excel_document(doc: pdf.Document, selected: list[int], output: Path, config: dict) -> int:
    from openpyxl import Workbook
    from openpyxl.drawing.image import Image
    from openpyxl.styles import Border, PatternFill, Side
    from openpyxl.utils import get_column_letter
    from pdf_engine import excel_text, render

    workbook = Workbook()
    workbook.remove(workbook.active)
    fallback = 0
    for number in selected:
        page = doc[number]
        layout = page_layout(page, config)
        sheet = workbook.create_sheet(f"Pagina {number + 1}")
        sheet.sheet_view.showGridLines = False
        reference = workbook.create_sheet(f"Original {number + 1}")
        image = Image(io.BytesIO(render(page, 144, config).tobytes("png")))
        image.width, image.height = page.rect.width * 96 / 72, page.rect.height * 96 / 72
        reference.add_image(image, "A1")
        if layout["fallback"]:
            sheet["A1"] = "Pagina conservada como imagen en la hoja Original; use OCR para recuperar texto."
            sheet.column_dimensions["A"].width = 90
            fallback += 1
            continue
        row_index = 1
        drawings = page.get_drawings()
        columns = max((len(table["xs"]) - 1 for table in layout["tables"]), default=1)
        events = [(table["bbox"], "table", table) for table in layout["tables"]]
        for line in layout["lines"]:
            spans = [span for span in line["spans"]
                     if not any(contains(table["bbox"], span["bbox"]) for table in layout["tables"])]
            if spans:
                events.append((line["bbox"], "text", {**line, "spans": spans}))
        for _, kind, content in sorted(events, key=lambda event: (event[0][1], event[0][0])):
            if kind == "text":
                spans = content["spans"]
                cell = sheet.cell(row_index, 1, excel_text(" ".join(span["text"] for span in spans)))
                style_excel_cell(cell, spans)
                cell.fill = PatternFill("solid", fgColor=cell_fill(drawings, content["bbox"]))
                sheet.row_dimensions[row_index].height = max(span["size"] for span in spans) * 1.5
                if columns > 1:
                    sheet.merge_cells(start_row=row_index, start_column=1,
                                      end_row=row_index, end_column=columns)
                else:
                    sheet.column_dimensions["A"].width = 80
                row_index += 1
            else:
                xs, ys = content["xs"], content["ys"]
                for column, width in enumerate([b - a for a, b in zip(xs, xs[1:])], 1):
                    letter = get_column_letter(column)
                    sheet.column_dimensions[letter].width = max(
                        sheet.column_dimensions[letter].width, min(100, width * 96 / 72 / 7)
                    )
                for row, column, end_row, end_col, rect in content["cells"]:
                    cell_lines = lines_in(layout["lines"], rect)
                    text = "\n".join(" ".join(span["text"] for span in line["spans"]) for line in cell_lines)
                    cell = sheet.cell(row_index + row, column + 1, excel_text(text))
                    spans = [span for line in cell_lines for span in line["spans"]]
                    style_excel_cell(cell, spans)
                    cell.fill = PatternFill("solid", fgColor=cell_fill(drawings, rect))
                    cell.border = Border(*(Side(style="thin", color="808080") for _ in range(4)))
                    if row != end_row or column != end_col:
                        sheet.merge_cells(start_row=row_index + row, start_column=column + 1,
                                          end_row=row_index + end_row, end_column=end_col + 1)
                for row, height in enumerate([b - a for a, b in zip(ys, ys[1:])]):
                    sheet.row_dimensions[row_index + row].height = height
                row_index += len(ys)
        sheet.freeze_panes = "A2"
    workbook.save(output)
    workbook.close()
    return fallback


def style_excel_cell(cell, spans: list[dict]) -> None:
    from openpyxl.styles import Alignment, Font

    cell.alignment = Alignment(vertical="top", wrap_text=True)
    if spans:
        span = spans[0]
        cell.font = Font(name=font_name(span), size=span["size"], bold=bool(span["flags"] & 16),
                         italic=bool(span["flags"] & 2), color=rgb(span))


def ppt_runs(frame, lines: list[dict], scale: float) -> None:
    from pptx.dml.color import RGBColor
    from pptx.util import Pt
    from pdf_engine import clean_text

    frame.clear()
    frame.margin_left = frame.margin_right = frame.margin_top = frame.margin_bottom = 0
    frame.word_wrap = False
    for index, line in enumerate(lines):
        paragraph = frame.paragraphs[0] if not index else frame.add_paragraph()
        paragraph.space_before = paragraph.space_after = Pt(0)
        paragraph.line_spacing = 1.0
        previous = None
        for span in line["spans"]:
            run = paragraph.add_run()
            separator = " " if previous is not None and span["bbox"][0] - previous["bbox"][2] > 1 else ""
            run.text = separator + clean_text(span["text"])
            run.font.name = font_name(span)
            run.font.size = Pt(max(1, span["size"] * scale))
            run.font.bold = bool(span["flags"] & 16)
            run.font.italic = bool(span["flags"] & 2)
            run.font.color.rgb = RGBColor.from_string(rgb(span))
            previous = span


def powerpoint_document(doc: pdf.Document, selected: list[int], output: Path, config: dict) -> int:
    from pptx import Presentation
    from pptx.dml.color import RGBColor
    from pptx.util import Inches
    from pdf_engine import render

    presentation = Presentation()
    first = doc[selected[0]].rect
    presentation.slide_width = Inches(10)
    presentation.slide_height = Inches(max(1, min(56, 10 * first.height / first.width)))
    fallback = 0
    for number in selected:
        page = doc[number]
        layout = page_layout(page, config)
        slide = presentation.slides.add_slide(presentation.slide_layouts[6])
        width, height = int(presentation.slide_width), int(presentation.slide_height)
        ratio = min(width / page.rect.width, height / page.rect.height)
        left, top = (width - int(page.rect.width * ratio)) // 2, (height - int(page.rect.height * ratio)) // 2
        if layout["fallback"]:
            image = layout["image"]
            fallback += 1
            lines, tables = [], []
        else:
            drawings = page.get_drawings()
            tables = [table for table in layout["tables"] if not any(
                tuple(line["dir"]) != (1.0, 0.0) and contains(table["bbox"], line["bbox"])
                for line in layout["lines"]
            )]
            lines = [line for line in layout["lines"] if tuple(line["dir"]) == (1.0, 0.0)]
            with pdf.open() as background:
                background.insert_pdf(doc, from_page=number, to_page=number)
                backdrop = background[0]
                for line in lines:
                    for span in line["spans"]:
                        backdrop.add_redact_annot(span["bbox"], fill=False)
                backdrop.apply_redactions(images=0, graphics=0)
                image = render(backdrop, 144, config).tobytes("png")
        slide.shapes.add_picture(io.BytesIO(image), left, top,
                                 width=int(page.rect.width * ratio), height=int(page.rect.height * ratio))
        for table in tables:
            xs, ys = table["xs"], table["ys"]
            shape = slide.shapes.add_table(
                len(ys) - 1, len(xs) - 1, left + int(xs[0] * ratio), top + int(ys[0] * ratio),
                max(1, int((xs[-1] - xs[0]) * ratio)), max(1, int((ys[-1] - ys[0]) * ratio)),
            )
            for column, size in zip(shape.table.columns, [b - a for a, b in zip(xs, xs[1:])]):
                column.width = max(1, int(size * ratio))
            for row, size in zip(shape.table.rows, [b - a for a, b in zip(ys, ys[1:])]):
                row.height = max(1, int(size * ratio))
            for row, column, end_row, end_col, rect in table["cells"]:
                cell = shape.table.cell(row, column)
                if row != end_row or column != end_col:
                    cell.merge(shape.table.cell(end_row, end_col))
                cell.fill.solid()
                cell.fill.fore_color.rgb = RGBColor.from_string(cell_fill(drawings, rect))
                ppt_runs(cell.text_frame, lines_in(lines, rect), ratio / 12700)
        for line in lines:
            spans = [span for span in line["spans"]
                     if not any(contains(table["bbox"], span["bbox"]) for table in tables)]
            for span in spans:
                x0, y0, x1, y1 = span["bbox"]
                shape = slide.shapes.add_textbox(
                    left + int(x0 * ratio), top + int(y0 * ratio),
                    max(1, int((x1 - x0 + 2) * ratio)), max(1, int((y1 - y0 + 2) * ratio)),
                )
                ppt_runs(shape.text_frame, [{**line, "spans": [span]}], ratio / 12700)
    presentation.save(output)
    return fallback


def export_office(operation: str, doc: pdf.Document, selected: list[int],
                  folder: Path, config: dict) -> tuple[Path, str]:
    converters = {
        "pdf_to_docx": (word_document, "documento.docx", "Word: texto y tablas editables en su posicion original sobre un fondo con logos, firmas, marcas y bordes. Disposicion fija: textos largos nuevos no reacomodan el resto de la pagina."),
        "pdf_to_xlsx": (excel_document, "tablas.xlsx", "Excel: tablas editables y texto; cada hoja Original conserva la pagina como imagen. Sin formulas inferidas."),
        "pdf_to_pptx": (powerpoint_document, "presentacion.pptx", "PowerPoint: texto horizontal y tablas editables; imagenes, trazos y texto girado en fondo no editable."),
    }
    convert, filename, notice = converters[operation]
    output = folder / filename
    fallback = convert(doc, selected, output, config)
    if fallback:
        notice += f" {fallback} paginas complejas, escaneadas/sin texto visible o giradas conservadas como imagen; OCR recupera texto, no reconstruye el diseno."
    notice += " Fuentes sustituidas si no estan instaladas; revise columnas, saltos, colores y celdas combinadas. No equivale a Acrobat."
    return output, notice
