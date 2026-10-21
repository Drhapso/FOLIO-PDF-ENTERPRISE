from __future__ import annotations

import base64
import io
import json
import math
import os
import re
import secrets
import shutil
import subprocess
import zipfile
from contextlib import ExitStack
from pathlib import Path
from xml.etree import ElementTree

import pymupdf as pdf
from PIL import Image, ImageOps

OPERATIONS = {
    "merge", "split", "organize", "rotate", "compress", "images_to_pdf",
    "pdf_to_images", "extract_text", "ocr", "edit", "watermark",
    "protect", "unlock", "pdf_to_docx", "pdf_to_xlsx", "pdf_to_pptx",
    "office_to_pdf", "inspect", "edit_text_pro",
    "extract_pages", "delete_pages", "number_pages", "extract_images", "crop",
    "resize", "nup", "overlay", "compare", "remove_metadata", "grayscale",
    "rasterize", "flatten", "repair", "web_optimize", "blank_report",
    "remove_blank", "form_fields", "fill_form", "dossier", "stamp",
    "html_to_pdf", "pdf_to_svg",
    "heic_to_jpg", "epub_to_pdf", "pdf_to_epub", "verify_signatures",
}
IMAGE_TYPES = {".png", ".jpg", ".jpeg", ".tif", ".tiff", ".webp", ".bmp", ".heic", ".heif"}
OFFICE_TYPES = {".docx", ".xlsx", ".pptx", ".odt", ".ods", ".odp"}
# Combinan muchas entradas en un solo resultado.
COMBINE_OPERATIONS = {"merge", "images_to_pdf", "dossier"}
# Procesan cada PDF por separado y entregan un ZIP cuando hay varios.
BATCH_OPERATIONS = {
    "split", "compress", "pdf_to_images", "extract_text", "pdf_to_docx",
    "pdf_to_xlsx", "pdf_to_pptx", "pdf_to_svg", "extract_images",
    "pdf_to_epub",
}
BULK_OPERATIONS = COMBINE_OPERATIONS | BATCH_OPERATIONS


def file_limit(operation: str, config: dict) -> int:
    return config["max_batch_files"] if operation in BULK_OPERATIONS else config["max_files"]


class ProcessingError(Exception):
    def __init__(self, message: str, status: int = 422):
        super().__init__(message)
        self.status = status


def libreoffice_path(config: dict) -> str | None:
    configured = config["libreoffice_path"]
    if configured:
        return configured if Path(configured).is_file() else None
    found = shutil.which("soffice")
    if found:
        return found
    for variable in ("PROGRAMFILES", "PROGRAMFILES(X86)"):
        folder = os.environ.get(variable)
        if folder:
            candidate = Path(folder) / "LibreOffice" / "program" / "soffice.com"
            if candidate.is_file():
                return str(candidate)
    return None


def capabilities(config: dict, root: Path) -> dict:
    languages = config["ocr_language"].split("+")
    return {
        "office": bool(libreoffice_path(config)),
        "ocr": all((root / "models" / f"{lang}.traineddata").is_file() for lang in languages),
        "ocr_language": config["ocr_language"],
    }


def integer(value, low: int, high: int, label: str) -> int:
    if isinstance(value, bool) or not re.fullmatch(r"-?\d+", str(value)):
        raise ProcessingError(f"{label}: debe ser un entero.")
    number = int(value)
    if not low <= number <= high:
        raise ProcessingError(f"{label}: valor fuera de rango ({low}-{high}).")
    return number


def pages_from(spec: str, count: int, *, repeats: bool = False) -> list[int]:
    if not isinstance(spec, str) or len(spec) > 10000:
        raise ProcessingError("Lista de paginas no valida.")
    if not spec.strip():
        return list(range(count))
    result = []
    for token in spec.split(","):
        match = re.fullmatch(r"\s*(\d+)(?:\s*-\s*(\d+))?\s*", token)
        if not match:
            raise ProcessingError("Use paginas como 1,3,5-8.")
        start = integer(match[1], 1, count, "Pagina")
        end = integer(match[2] or match[1], 1, count, "Pagina")
        step = 1 if end >= start else -1
        result.extend(range(start - 1, end - 1 + step, step))
        if len(result) > count * 2:
            raise ProcessingError("Demasiadas paginas seleccionadas.")
    if not repeats and len(result) != len(set(result)):
        raise ProcessingError("No repita paginas en la seleccion.")
    return result


def open_pdf(path: Path, password: str, limit: int) -> pdf.Document:
    document = pdf.open(path)
    try:
        if not document.is_pdf:
            raise ProcessingError("El archivo no es un PDF.")
        if document.needs_pass and not document.authenticate(password):
            raise ProcessingError("El PDF esta protegido. Introduzca la clave correcta.")
        if not len(document):
            raise ProcessingError("El documento no contiene paginas.")
        if limit and len(document) > limit:
            raise ProcessingError(f"El documento supera el limite configurado de {limit} paginas.")
        return document
    except BaseException:
        document.close()
        raise


def render(page: pdf.Page, dpi: int, config: dict) -> pdf.Pixmap:
    pixels = math.ceil(page.rect.width * dpi / 72) * math.ceil(page.rect.height * dpi / 72)
    if pixels > config["max_render_pixels"]:
        raise ProcessingError("Pagina demasiado grande para esa resolucion. Reduzca los DPI.")
    return page.get_pixmap(dpi=dpi, colorspace=pdf.csRGB, alpha=False)


def save(document: pdf.Document, path: Path, **kwargs):
    document.save(path, garbage=4, deflate=True, use_objstms=1, **kwargs)


def inspect_document(doc: pdf.Document, options: dict, config: dict) -> dict:
    from acrobat_tools import extract_page_blocks, text_background
    from scan_editor import editable_blocks

    number = integer(options.get("page", 1), 1, len(doc), "Pagina")
    page = doc[number - 1]
    dpi = integer(options.get("dpi", 90), 72, 300, "DPI") if "dpi" in options else 90
    editable_background = options.get("text_background") is True
    if editable_background and page.rotation:
        raise ProcessingError("Quite la rotacion con Rotar PDF antes de editar el texto.")
    pix = render(page, dpi, config)
    payload = {
        "pages": len(doc), "page": number, "width": page.rect.width,
        "height": page.rect.height, "dpi": dpi,
        "image": base64.b64encode(
            pix.tobytes("png") if editable_background else pix.tobytes("jpeg", jpg_quality=80)
        ).decode("ascii"),
        "blocks": editable_blocks(page, config) if editable_background else extract_page_blocks(page),
    }
    if editable_background:
        payload["background"] = text_background(doc, number - 1, dpi, config, payload["blocks"])
        payload["scan_notice"] = (
            "Escaneo: reconstruccion local conservadora. OCR automatico; revise el texto. "
            "Las zonas complejas estan bloqueadas y no se modificaran."
            if any(block.get("source") == "scan" for block in payload["blocks"]) else ""
        )
        payload["image_mime"] = "image/png"
    return payload


def validate_office(path: Path):
    with zipfile.ZipFile(path) as archive:
        entries = archive.infolist()
        if len(entries) > 10000 or sum(entry.file_size for entry in entries) > 200 * 1024 * 1024:
            raise ProcessingError("El documento Office descomprimido excede el limite.")
        for entry in entries:
            name = entry.filename.lower()
            if any(part in name for part in ("vbaproject", "/embeddings/", "/activex/")):
                raise ProcessingError("No se admiten macros, ActiveX ni objetos incrustados.")
            if name.endswith(".rels"):
                root = ElementTree.fromstring(archive.read(entry))
                if any(child.get("TargetMode") == "External" for child in root):
                    raise ProcessingError("Quite los vinculos y recursos externos antes de convertir.")
            if name == "content.xml":
                root = ElementTree.fromstring(archive.read(entry))
                for node in root.iter():
                    for attribute, target in node.attrib.items():
                        if attribute.endswith("}href") and not target.startswith(("#", "Pictures/", "./Pictures/")):
                            raise ProcessingError("Quite los vinculos externos del documento OpenDocument.")


def convert_office(source: Path, folder: Path, config: dict) -> Path:
    executable = libreoffice_path(config)
    if not executable:
        raise ProcessingError("LibreOffice no esta instalado o su ruta no es valida.", 503)
    validate_office(source)
    profile = folder / "office-profile"
    user = profile / "user"
    user.mkdir(parents=True)
    (user / "registrymodifications.xcu").write_text(
        '<?xml version="1.0"?><oor:items xmlns:oor="http://openoffice.org/2001/registry">'
        '<item oor:path="/org.openoffice.Office.Common/Security/Scripting">'
        '<prop oor:name="MacroSecurityLevel" oor:op="fuse"><value>3</value></prop></item>'
        '<item oor:path="/org.openoffice.Office.Common/Save/Document">'
        '<prop oor:name="AutoSave" oor:op="fuse"><value>false</value></prop></item>'
        '</oor:items>', encoding="utf-8",
    )
    target = folder / "office-output"
    target.mkdir()
    completed = subprocess.run(
        [executable, f"-env:UserInstallation={profile.as_uri()}", "--headless",
         "--nologo", "--nodefault", "--norestore", "--convert-to", "pdf",
         "--outdir", str(target), str(source)],
        stdout=subprocess.DEVNULL, stderr=subprocess.DEVNULL,
        timeout=max(5, config["timeout_seconds"] - 5), check=False,
    )
    result = target / (source.stem + ".pdf")
    if completed.returncode or not result.is_file():
        raise ProcessingError("LibreOffice no pudo convertir el documento.")
    return result


def validate_options(options: dict):
    if not isinstance(options, dict):
        raise ProcessingError("Opciones no validas.")
    for key in ("password", "new_password", "pages", "groups", "text", "compression", "preprocess"):
        if key in options and (not isinstance(options[key], str) or len(options[key]) > 10000):
            raise ProcessingError(f"Opcion no valida: {key}.")
    for key in ("password", "new_password"):
        if len(options.get(key, "").encode("utf-8")) > 40:
            raise ProcessingError("La clave no puede superar 40 bytes UTF-8.")


def process(operation: str, sources: list[Path], options: dict, folder: Path, config: dict, root: Path) -> dict:
    if operation not in OPERATIONS:
        raise ProcessingError("Herramienta desconocida.", 404)
    validate_options(options)
    if not sources or len(sources) > file_limit(operation, config):
        raise ProcessingError("Cantidad de archivos no valida.")
    if operation in BATCH_OPERATIONS and len(sources) > 1:
        return process_batch(operation, sources, options, folder, config, root)
    from acrobat_tools import ACROBAT_OPERATIONS, process_acrobat
    if operation in ACROBAT_OPERATIONS:
        return process_acrobat(operation, sources, options, folder, config)
    from advanced_tools import ADVANCED_OPERATIONS, process_advanced
    if operation in ADVANCED_OPERATIONS:
        return process_advanced(operation, sources, options, folder, config)
    from extended_tools import EXTENDED_OPERATIONS, process_extended
    if operation in EXTENDED_OPERATIONS:
        return process_extended(operation, sources, options, folder, config)
    if operation not in COMBINE_OPERATIONS and len(sources) != 1:
        raise ProcessingError("Esta herramienta admite un solo archivo.")
    if operation == "merge" and len(sources) < 2:
        raise ProcessingError("Seleccione al menos dos PDF.")
    allowed = IMAGE_TYPES if operation == "images_to_pdf" else OFFICE_TYPES if operation == "office_to_pdf" else {".pdf"}
    if any(source.suffix.lower() not in allowed for source in sources):
        raise ProcessingError("Formato no admitido para esta herramienta.")
    output = folder / "resultado.pdf"
    notice = ""
    if operation == "office_to_pdf":
        converted = convert_office(sources[0], folder, config)
        with open_pdf(converted, "", config["max_pages"]):
            pass
        shutil.move(converted, output)
    elif operation == "images_to_pdf":
        from pillow_heif import register_heif_opener
        register_heif_opener()
        with pdf.open() as combined:
            for source in sources:
                with Image.open(source) as image:
                    if image.width * image.height > config["max_render_pixels"]:
                        raise ProcessingError("Imagen demasiado grande.")
                    if getattr(image, "n_frames", 1) != 1:
                        raise ProcessingError("Use imagenes de un solo fotograma.")
                    image = ImageOps.exif_transpose(image)
                    if image.mode in ("RGBA", "LA") or "transparency" in image.info:
                        rgba = image.convert("RGBA")
                        flat = Image.new("RGB", rgba.size, "white")
                        flat.paste(rgba, mask=rgba.getchannel("A"))
                        image = flat
                    image = image.convert("RGB")
                    data = io.BytesIO()
                    image.save(data, format="JPEG", quality=95)
                    page = combined.new_page(width=image.width * 0.75, height=image.height * 0.75)
                    page.insert_image(page.rect, stream=data.getvalue())
            if config["max_batch_pages"] and len(combined) > config["max_batch_pages"]:
                raise ProcessingError("Demasiadas paginas.")
            save(combined, output)
    else:
        with ExitStack() as stack:
            documents = [
                stack.enter_context(open_pdf(source, options.get("password", ""), config["max_pages"]))
                for source in sources
            ]
            total_limit = config["max_batch_pages"] if operation == "merge" else config["max_pages"]
            if total_limit and sum(map(len, documents)) > total_limit:
                raise ProcessingError("La suma de paginas excede el limite configurado.")
            doc = documents[0]
            selected = pages_from(options.get("pages", ""), len(doc))
            if operation == "inspect":
                output = folder / "vista.json"
                output.write_text(json.dumps(inspect_document(doc, options, config)), encoding="utf-8")
            elif operation == "merge":
                with pdf.open() as combined:
                    for document in documents:
                        combined.insert_pdf(document)
                    save(combined, output)
            elif operation == "split":
                spec = options.get("groups", "").strip()
                if spec and any(not group.strip() for group in spec.split(";")):
                    raise ProcessingError("No deje grupos de paginas vacios.")
                groups = [pages_from(group, len(doc)) for group in spec.split(";")] if spec else [[n] for n in selected]
                if config["max_pages"] and (len(groups) > config["max_pages"] or
                                           sum(map(len, groups)) > config["max_pages"] * 2):
                    raise ProcessingError("Demasiadas paginas de salida.")
                output = folder / "paginas.zip"
                with zipfile.ZipFile(output, "w", zipfile.ZIP_DEFLATED) as archive:
                    for index, group in enumerate(groups, 1):
                        with pdf.open() as part:
                            for number in group:
                                part.insert_pdf(doc, from_page=number, to_page=number)
                            archive.writestr(f"parte-{index:03}.pdf", part.tobytes(garbage=4, deflate=True))
                        check_size(output, config)
            elif operation == "organize":
                doc.select(selected)
                save(doc, output)
            elif operation == "rotate":
                angle = integer(options.get("angle", 90), 90, 270, "Angulo")
                if angle % 90:
                    raise ProcessingError("El angulo debe ser 90, 180 o 270.")
                for number in selected:
                    doc[number].set_rotation((doc[number].rotation + angle) % 360)
                save(doc, output)
            elif operation == "compress":
                mode = options.get("compression", "lossless")
                if mode == "lossless":
                    save(doc, output)
                elif mode == "raster":
                    with pdf.open() as reduced:
                        dpi = integer(options.get("dpi", 120), 72, 200, "DPI")
                        for page in doc:
                            pix = render(page, dpi, config)
                            new = reduced.new_page(width=page.rect.width, height=page.rect.height)
                            new.insert_image(new.rect, stream=pix.tobytes("jpeg", jpg_quality=65))
                        save(reduced, output)
                    notice = "Compresion visual: elimina texto seleccionable, formularios, enlaces y capas."
                else:
                    raise ProcessingError("Modo de compresion desconocido.")
                if output.stat().st_size >= sources[0].stat().st_size:
                    shutil.copyfile(sources[0], output)
                    notice = "No se obtuvo reduccion; se entrega el original sin cambios."
            elif operation == "pdf_to_images":
                dpi = integer(options.get("dpi", 144), 72, 300, "DPI")
                output = folder / "imagenes.zip"
                image_format = options.get("image_format", "png")
                if image_format not in {"png", "jpg", "tiff"}:
                    raise ProcessingError("Formato de imagen no valido.")
                with zipfile.ZipFile(output, "w", zipfile.ZIP_DEFLATED) as archive:
                    for number in selected:
                        pix = render(doc[number], dpi, config)
                        if image_format == "tiff":
                            buffer = io.BytesIO()
                            Image.frombytes("RGB", (pix.width, pix.height), pix.samples).save(buffer, "TIFF", compression="tiff_deflate")
                            data = buffer.getvalue()
                        else:
                            data = pix.tobytes("jpeg" if image_format == "jpg" else "png")
                        archive.writestr(f"pagina-{number + 1:03}.{image_format}", data)
                        check_size(output, config)
            elif operation == "extract_text":
                output = folder / "texto.txt"
                output.write_text("\n\n".join(doc[n].get_text(sort=True) for n in selected), encoding="utf-8")
                if not output.read_text(encoding="utf-8").strip():
                    raise ProcessingError("No hay texto extraible. Ejecute OCR primero.")
            elif operation == "ocr":
                if not capabilities(config, root)["ocr"]:
                    raise ProcessingError("Faltan los modelos OCR locales. Ejecute deploy\\install-ocr.ps1.", 503)
                dpi = integer(options.get("dpi", 200), 100, 300, "DPI")
                from advanced_tools import preprocess_scan
                mode = options.get("preprocess", "none")
                with pdf.open() as searchable:
                    for page in doc:
                        pix = render(page, dpi, config)
                        pix = preprocess_scan(pix, mode, dpi, config["max_render_pixels"])
                        raw = pix.pdfocr_tobytes(language=config["ocr_language"], tessdata=str(root / "models"))
                        with pdf.open(stream=raw, filetype="pdf") as recognized:
                            searchable.insert_pdf(recognized)
                    save(searchable, output)
                notice = "OCR automatico: revise nombres, cifras y fechas; no garantiza transcripcion exacta."
                if mode != "none":
                    notice += " Enderezado hasta 5 grados; limpieza opcional en blanco y negro. Revise sellos y trazos tenues."
            elif operation == "watermark":
                text = options.get("text", "").strip()
                if not text or len(text) > 100:
                    raise ProcessingError("Escriba una marca de agua de 1 a 100 caracteres.")
                for number in selected:
                    page = doc[number]
                    result = page.insert_textbox(
                        pdf.Rect(20, page.rect.height / 2 - 30, page.rect.width - 20, page.rect.height / 2 + 60),
                        text, fontsize=24, color=(0.4, 0.4, 0.4), align=1, overlay=True, fill_opacity=0.35,
                    )
                    if result < 0:
                        raise ProcessingError("La marca de agua no cabe en una pagina.")
                save(doc, output)
            elif operation == "protect":
                password = options.get("new_password", "")
                if len(password) < 8:
                    raise ProcessingError("La nueva clave debe tener al menos 8 caracteres.")
                save(doc, output, encryption=pdf.PDF_ENCRYPT_AES_256,
                     owner_pw=secrets.token_hex(16), user_pw=password,
                     permissions=pdf.PDF_PERM_PRINT | pdf.PDF_PERM_COPY | pdf.PDF_PERM_ANNOTATE)
            elif operation == "unlock":
                if not options.get("password"):
                    raise ProcessingError("Debe introducir la clave actual.")
                save(doc, output, encryption=pdf.PDF_ENCRYPT_NONE)
            elif operation == "pdf_to_docx":
                from docx import Document
                from docx.shared import Inches

                result = Document()
                for index, number in enumerate(selected):
                    if index:
                        result.add_page_break()
                    text = doc[number].get_text(sort=True)
                    if text.strip():
                        for line in text.splitlines():
                            result.add_paragraph(clean_text(line))
                    else:
                        result.add_picture(io.BytesIO(render(doc[number], 144, config).tobytes("png")), width=Inches(6))
                output = folder / "documento.docx"
                result.save(output)
                notice = "Word: texto extraido sin reconstruccion exacta de diseno; paginas sin texto como imagen."
            elif operation == "pdf_to_xlsx":
                from openpyxl import Workbook

                workbook = Workbook()
                workbook.remove(workbook.active)
                for number in selected:
                    sheet = workbook.create_sheet(f"Pagina {number + 1}")
                    tables = doc[number].find_tables().tables
                    rows = []
                    for table in tables:
                        rows.extend(table.extract())
                        rows.append([])
                    if not tables:
                        rows = [[line] for line in doc[number].get_text(sort=True).splitlines()]
                    for row in rows:
                        sheet.append([excel_text(cell) for cell in row])
                output = folder / "tablas.xlsx"
                workbook.save(output)
                notice = "Excel: tablas detectadas automaticamente; si no hay tablas se extrae texto por linea."
            elif operation == "pdf_to_pptx":
                from pptx import Presentation
                from pptx.util import Inches

                presentation = Presentation()
                first = doc[selected[0]].rect
                presentation.slide_width = Inches(10)
                presentation.slide_height = Inches(max(1, min(56, 10 * first.height / first.width)))
                for number in selected:
                    page = doc[number]
                    slide = presentation.slides.add_slide(presentation.slide_layouts[6])
                    height = int(presentation.slide_height)
                    width = int(presentation.slide_width)
                    ratio = min(width / page.rect.width, height / page.rect.height)
                    w, h = int(page.rect.width * ratio), int(page.rect.height * ratio)
                    slide.shapes.add_picture(
                        io.BytesIO(render(page, 144, config).tobytes("png")),
                        (width - w) // 2, (height - h) // 2, width=w, height=h,
                    )
                output = folder / "presentacion.pptx"
                presentation.save(output)
                notice = "PowerPoint: cada pagina se conserva como imagen, no como objetos editables."
    check_size(output, config)
    return {"file": output.name, "notice": notice}


TEXT_RESULTS = {".txt", ".json", ".svg", ".csv"}


def process_batch(operation: str, sources: list[Path], options: dict, folder: Path, config: dict, root: Path) -> dict:
    """Procesa cada PDF en su propio subespacio y reune los resultados en un ZIP."""
    output = folder / "lote.zip"
    notices = []
    width = max(2, len(str(len(sources))))
    with zipfile.ZipFile(output, "w", zipfile.ZIP_STORED, allowZip64=True) as archive:
        for index, source in enumerate(sources, 1):
            label = f"{index:0{width}}"
            work = folder / f"lote-{label}"
            work.mkdir()
            try:
                try:
                    result = process(operation, [source], options, work, config, root)
                except ProcessingError as error:
                    raise ProcessingError(f"Archivo {index}: {error}", error.status) from None
                except (pdf.FileDataError, pdf.EmptyFileError):
                    raise ProcessingError(f"Archivo {index}: PDF danado o no procesable.", 422) from None
                produced = work / result["file"]
                if produced.suffix == ".zip":
                    with zipfile.ZipFile(produced) as inner:
                        for entry in inner.infolist():
                            if entry.is_dir():
                                continue
                            name = f"{label}/{Path(entry.filename).name}"
                            info = zipfile.ZipInfo(name, date_time=entry.date_time)
                            info.compress_type = zipfile.ZIP_DEFLATED if Path(name).suffix in TEXT_RESULTS else zipfile.ZIP_STORED
                            with inner.open(entry) as reader, archive.open(info, "w", force_zip64=True) as writer:
                                shutil.copyfileobj(reader, writer, 1024 * 1024)
                else:
                    name = f"{label}-{produced.name}"
                    kind = zipfile.ZIP_DEFLATED if produced.suffix in TEXT_RESULTS else zipfile.ZIP_STORED
                    archive.write(produced, name, compress_type=kind)
                if result["notice"] and result["notice"] not in notices:
                    notices.append(result["notice"])
            finally:
                shutil.rmtree(work, ignore_errors=True)
            # En error la carpeta del trabajo se elimina completa al terminar.
            source.unlink(missing_ok=True)
            check_size(output, config)
    notice = f"{len(sources)} archivos procesados por separado; el ZIP sigue el orden de seleccion (01, 02...)."
    return {"file": output.name, "notice": " ".join([notice, *notices])}


def clean_text(text: str) -> str:
    from openpyxl.cell.cell import ILLEGAL_CHARACTERS_RE

    return ILLEGAL_CHARACTERS_RE.sub("", text)


def excel_text(value) -> str:
    text = clean_text(str(value or ""))[:32767]
    return "'" + text if text.startswith(("=", "+", "-", "@")) else text


def check_size(path: Path, config: dict):
    if path.stat().st_size > config["max_output_mb"] * 1024 * 1024:
        raise ProcessingError("El resultado excede el limite de salida. Procese menos paginas.", 413)
